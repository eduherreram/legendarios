from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "funcionamiento-por-rol.md"
OUTPUT = ROOT / "docs" / "funcionamiento-por-rol.docx"

ACCENT = "1F4E5F"
MUTED = "64748B"
LIGHT = "EAF5F3"
BORDER = "BFD8D3"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, margin: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color, before, after in [
        ("Heading 1", 22, ACCENT, 18, 8),
        ("Heading 2", 16, ACCENT, 14, 6),
        ("Heading 3", 12, "334155", 10, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = kicker.add_run("DOCUMENTACION FUNCIONAL")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Legendario Manager")
    r.font.name = "Aptos Display"
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Funcionamiento del proyecto separado por rol")
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string("334155")

    doc.add_paragraph()
    note_table = doc.add_table(rows=1, cols=1)
    note_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = note_table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_border(cell, color="A9CFC8")
    set_cell_margins(cell, 180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Resumen de roles, permisos, modulos, flujos operativos y estructura tecnica del sistema."
    )
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("334155")

    for _ in range(9):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Generado desde docs/funcionamiento-por-rol.md")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("Indice del documento", level=1)
    items = [
        "Resumen general",
        "Roles del sistema",
        "Vista por rol",
        "Modulos principales",
        "Flujos importantes",
        "Matriz resumida de permisos",
        "Estructura tecnica",
        "Datos y relaciones principales",
        "Instalacion y ejecucion local",
        "Observaciones para mantenimiento",
    ]
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)
    doc.add_page_break()


def apply_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string("0F766E")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_border(cell, color="D9E2E8")
    set_cell_margins(cell, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("334155")


def add_single_table(doc: Document, rows: list[list[str]], compact: bool = False) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])

    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 90 if compact else 130)
            set_cell_border(cell, color=BORDER)
            if i == 0:
                set_cell_shading(cell, ACCENT)
            elif i % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (compact or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            run.font.size = Pt(7.4 if compact else 9.2)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
            else:
                run.font.color.rgb = RGBColor.from_string("334155")

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
    doc.add_paragraph()


def add_markdown_table(doc: Document, rows: list[list[str]], compact: bool = False) -> None:
    if rows and len(rows[0]) > 5:
        left_rows = [[row[0], *row[1:5]] for row in rows]
        right_rows = [[row[0], *row[5:]] for row in rows]

        caption = doc.add_paragraph()
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(4)
        run = caption.add_run("Permisos operativos principales")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("334155")
        add_single_table(doc, left_rows, compact=True)

        doc.add_page_break()
        caption = doc.add_paragraph()
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(4)
        run = caption.add_run("Permisos complementarios y roles publicos")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("334155")
        add_single_table(doc, right_rows, compact=True)
        return

    add_single_table(doc, rows, compact=compact)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip()
        if re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", raw):
            index += 1
            continue
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        rows.append(cells)
        index += 1
    return rows, index


def add_content(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    in_code = False
    code_lines: list[str] = []
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith("|"):
            rows, next_index = parse_table(lines, i)
            add_markdown_table(doc, rows, compact=len(rows[0]) > 4)
            i = next_index
            continue

        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            apply_inline_runs(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            text = re.sub(r"^\d+\. ", "", line)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            apply_inline_runs(p, text)
        else:
            p = doc.add_paragraph()
            apply_inline_runs(p, line)
        i += 1


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc)
    add_content(doc, markdown)
    doc.core_properties.title = "Funcionamiento del Proyecto por Rol"
    doc.core_properties.subject = "Legendario Manager"
    doc.core_properties.author = "Codex"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
